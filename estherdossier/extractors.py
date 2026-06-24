"""Tekst-extractie per bestandstype.

Elke extractor geeft platte tekst terug. Onbekende of niet-ondersteunde
typen geven een lege string + een notitie, zodat het altijd traceerbaar is
welke bestanden niet (volledig) uitgelezen konden worden.
"""
from __future__ import annotations

from pathlib import Path

SUPPORTED = {".docx", ".pdf", ".xlsx", ".pptx", ".rtf", ".txt", ".csv", ".md", ".log"}


def extract_text(path: Path) -> tuple[str, str]:
    """Geef (tekst, notitie) terug. notitie is leeg als alles goed ging."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            return _docx(path), ""
        if suffix == ".pdf":
            return _pdf(path)
        if suffix == ".xlsx":
            return _xlsx(path), ""
        if suffix == ".pptx":
            return _pptx(path), ""
        if suffix == ".rtf":
            return _rtf(path), ""
        if suffix in {".txt", ".csv", ".md", ".log"}:
            return path.read_text(encoding="utf-8", errors="replace"), ""
        return "", f"Niet-ondersteund bestandstype: {suffix or '(geen extensie)'}"
    except Exception as exc:  # robuust: één kapot bestand stopt de run niet
        return "", f"Fout bij uitlezen ({type(exc).__name__}): {exc}"


def _docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf(path: Path) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    text = "\n\n".join(pages).strip()
    note = ""
    if not text:
        note = "PDF bevat geen extraheerbare tekst (waarschijnlijk gescand beeld; OCR nodig)."
    return text, note


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"### Werkblad: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"### Dia {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
    return "\n".join(parts)


def _rtf(path: Path) -> str:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(path.read_text(encoding="utf-8", errors="replace"))
